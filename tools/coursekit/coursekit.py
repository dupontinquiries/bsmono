#!/usr/bin/env python3
"""coursekit -- turn a folder of lecture files into notes you can actually use.

    coursekit.py convert <in> [<in>...] -o notes/     one file or a whole tree
    coursekit.py probe   <in>                         what would happen, no writes

Handles .docx, .pptx, .pdf. Emits, per source file:

    notes/<Name>.md
    notes/<Name>.assets/<image files>

That sidecar layout is the Typora / Mark Text convention, and it is what
voiceterm already resolves -- `manpages::resolve_one_img_src` resolves a
relative img src against the note's directory. It is also the layout the
NUT1970 `md/` folder already uses, so these drop in beside existing notes.

Three things here that a plain `pandoc -o out.md` does not do, each fixing a
failure seen on real coursework:

* **Headings inferred from run size, not boldness.** Word files routinely
  define Heading1..6 and apply none of them -- the Elementals package used
  ListParagraph and nothing else. Every heading was bold direct formatting,
  so pandoc emitted bold paragraphs and the document had no outline at all.
  This reads `w:sz` off each paragraph's runs, clusters the distinct sizes,
  and maps the largest few onto h1/h2/h3.

* **Images deduplicated by content hash.** A 40-slide deck carries the
  faculty logo 40 times and the section divider 8 times. Naive extraction
  writes 40 files; this writes one and links it 40 times, or drops it
  entirely with --drop-repeats since a logo on every slide is furniture.

* **Slides keep their speaker notes.** The notes are usually where the
  lecturer said the thing that is actually examinable, and every
  slides-to-text path throws them away.
"""
import argparse, hashlib, os, re, shutil, subprocess, sys, zipfile
from collections import Counter

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# --------------------------------------------------------------- model

class Doc:
    def __init__(self, title):
        self.title = title
        self.blocks = []          # ('h',lvl,text) ('p',text) ('img',key) ('note',text)
        self.images = {}          # key -> bytes
        self.stats = Counter()

    def add_image(self, data, ext):
        """Content-addressed: identical bytes collapse to one key."""
        h = hashlib.sha256(data).hexdigest()[:12]
        key = f"{h}{ext}"
        if key in self.images:
            self.stats["image_repeats"] += 1
        else:
            self.images[key] = data
        return key


def slug(name):
    s = re.sub(r"[^\w\s.-]", "", name, flags=re.U).strip()
    return re.sub(r"\s+", " ", s)


# --------------------------------------------------------------- docx

def docx_doc(path):
    """Read paragraphs with their run sizes so headings can be inferred."""
    import docx
    from docx.shared import Pt

    d = docx.Document(path)
    doc = Doc(os.path.splitext(os.path.basename(path))[0])

    # media, indexed by relationship id
    rels = {}
    for rid, rel in d.part.rels.items():
        if "image" in rel.reltype:
            try:
                blob = rel.target_part.blob
                ext = os.path.splitext(rel.target_part.partname)[1] or ".png"
                rels[rid] = (blob, ext)
            except Exception:
                pass

    # pass 1: collect (text, size, bold, style) for every paragraph
    rows = []
    for p in d.paragraphs:
        txt = p.text.strip()
        sizes, bolds = [], []
        for r in p.runs:
            sz = r.font.size
            if sz is not None:
                sizes.append(sz.pt)
            bolds.append(bool(r.bold))
        style = (p.style.name or "").lower() if p.style else ""
        rows.append({
            "text": txt,
            "size": max(sizes) if sizes else None,
            "bold": bool(bolds) and all(bolds),
            "style": style,
            "para": p,
        })

    # pass 2: work out which sizes mean "heading"
    body = Counter(r["size"] for r in rows if r["text"] and r["size"])
    base = body.most_common(1)[0][0] if body else None
    big = sorted({s for s in body if base and s > base + 0.5}, reverse=True)[:3]
    level_of = {s: i + 1 for i, s in enumerate(big)}
    doc.stats["body_pt"] = base or 0
    doc.stats["heading_sizes"] = len(big)

    for r in rows:
        p = r["para"]
        # inline images first, in document order
        for blip in p._p.findall(f".//{{http://schemas.openxmlformats.org/drawingml/2006/main}}blip"):
            rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if rid in rels:
                blob, ext = rels[rid]
                doc.blocks.append(("img", doc.add_image(blob, ext)))
        if not r["text"]:
            continue
        lvl = None
        if r["style"].startswith("heading"):
            m = re.search(r"(\d)", r["style"])
            lvl = int(m.group(1)) if m else 1
            doc.stats["heading_from_style"] += 1
        elif r["size"] in level_of:
            lvl = level_of[r["size"]]
            doc.stats["heading_from_size"] += 1
        elif r["bold"] and len(r["text"]) < 90 and not r["text"].endswith((".", ",", ";")):
            lvl = 2
            doc.stats["heading_from_bold"] += 1
        if lvl:
            doc.blocks.append(("h", min(lvl, 6), r["text"]))
        else:
            doc.blocks.append(("p", r["text"]))

    for t in d.tables:
        doc.stats["tables_flattened"] += 1
        for row in t.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                doc.blocks.append(("p", " | ".join(cells)))
    return doc


# --------------------------------------------------------------- pptx

def pptx_doc(path):
    from pptx import Presentation

    pres = Presentation(path)
    doc = Doc(os.path.splitext(os.path.basename(path))[0])

    for n, slide in enumerate(pres.slides, 1):
        title = None
        body, pics = [], []
        for sh in slide.shapes:
            if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
                try:
                    img = sh.image
                    pics.append(doc.add_image(img.blob, "." + (img.ext or "png")))
                except Exception:
                    pass
                continue
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text.strip()
            if not txt:
                continue
            is_title = (slide.shapes.title is not None and sh == slide.shapes.title)
            if is_title and title is None:
                title = txt
            else:
                body.extend(l.strip() for l in txt.splitlines() if l.strip())

        doc.blocks.append(("h", 2, title or f"Slide {n}"))
        for k in pics:
            doc.blocks.append(("img", k))
        for line in body:
            doc.blocks.append(("p", line))

        # speaker notes: usually where the examinable remark lives
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text.strip()
            if nt:
                doc.stats["slides_with_notes"] += 1
                doc.blocks.append(("note", nt))
        doc.stats["slides"] += 1
    return doc


# ---------------------------------------------------------------- pdf

def pdf_doc(path, min_px=64):
    doc = Doc(os.path.splitext(os.path.basename(path))[0])
    txt = subprocess.run(["pdftotext", "-layout", path, "-"],
                         capture_output=True, text=True, errors="replace")
    if txt.returncode != 0:
        doc.stats["pdftotext_failed"] = 1
        return doc
    pages = txt.stdout.split("\f")

    tmp = f"/tmp/ck_{os.getpid()}"
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["pdfimages", "-png", "-p", path, f"{tmp}/i"],
                   capture_output=True)
    by_page = {}
    for f in sorted(os.listdir(tmp)):
        m = re.match(r"i-(\d+)-\d+\.png", f)
        if m:
            by_page.setdefault(int(m.group(1)), []).append(os.path.join(tmp, f))

    for n, page in enumerate(pages, 1):
        page = page.strip()
        if not page and n not in by_page:
            continue
        doc.blocks.append(("h", 2, f"Page {n}"))
        for f in by_page.get(n, []):
            # Filter on PIXEL DIMENSIONS, not file size: a flat-colour chart
            # compresses to a few hundred bytes and a byte threshold throws it
            # away. What we actually want to drop are rules and hairlines.
            try:
                from PIL import Image as _I
                with _I.open(f) as _im:
                    w_, h_ = _im.size
            except Exception:
                w_ = h_ = 0
            if w_ < min_px or h_ < min_px:
                doc.stats["slivers_skipped"] += 1
                continue
            doc.blocks.append(("img", doc.add_image(open(f, "rb").read(), ".png")))
        for para in re.split(r"\n\s*\n", page):
            para = " ".join(l.strip() for l in para.splitlines() if l.strip())
            if para:
                doc.blocks.append(("p", para))
        doc.stats["pages"] += 1
    shutil.rmtree(tmp, ignore_errors=True)
    return doc


# --------------------------------------------------------------- emit

def reencode(data, mode):
    """Re-encode one figure. Returns (bytes, extension).

    The markdown is byte-identical WITHIN A FAMILY, because the link is a
    relative path and only the bytes behind it change. Re-encoding to WebP
    renames the file, so there are two families:

        full  <-> none     (.png links)
        webp  <-> thumb    (.webp links)

    Which makes the useful move: ship `none` (text only, links intact), then
    drop the `full` assets folder in later and nothing needs rewriting. Or
    ship `thumb` and upgrade to `webp` the same way. Crossing families means
    re-running convert, which is cheap but is not a drop-in.
    """
    if mode == "full":
        return data, None
    from PIL import Image
    import io
    try:
        im = Image.open(io.BytesIO(data))
    except Exception:
        return data, None
    if im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    if mode == "thumb":
        im.thumbnail((480, 480))
        q = 70
    else:
        q = 82
    buf = io.BytesIO()
    im.save(buf, "WEBP", quality=q, method=6)
    return buf.getvalue(), ".webp"


def emit(doc, outdir, drop_repeats=0, link_titles=(), images="full"):
    os.makedirs(outdir, exist_ok=True)
    stem = slug(doc.title)
    assets = f"{stem}.assets"
    used = Counter(k for kind, *rest in
                   ((b[0], *b[1:]) for b in doc.blocks) if kind == "img"
                   for k in [rest[0]])
    furniture = {k for k, c in used.items() if drop_repeats and c >= drop_repeats}
    doc.stats["furniture_dropped"] = len(furniture)

    need = {k for k in used if k not in furniture}
    remap = {}
    if need and images != "none":
        os.makedirs(os.path.join(outdir, assets), exist_ok=True)
    for k in need:
        if images == "none":
            remap[k] = k          # link stays; the file simply is not there yet
            continue
        data, newext = reencode(doc.images[k], images)
        name = (os.path.splitext(k)[0] + newext) if newext else k
        remap[k] = name
        doc.stats["image_bytes"] += len(data)
        with open(os.path.join(outdir, assets, name), "wb") as f:
            f.write(data)

    out = [f"# {doc.title}", ""]
    for b in doc.blocks:
        if b[0] == "h":
            out += ["", "#" * b[1] + " " + b[2], ""]
        elif b[0] == "p":
            txt = b[1]
            for t in link_titles:
                if t != doc.title and len(t) > 7 and t in txt:
                    txt = txt.replace(t, f"[[{t}]]", 1)
            out.append(txt)
            out.append("")
        elif b[0] == "li":
            out.append(f"- {b[1]}")
        elif b[0] == "tablerow":
            out.append("| " + " | ".join(b[1]) + " |")
        elif b[0] == "tablesep":
            out.append("|" + "|".join("---" for _ in b[1]) + "|")
        elif b[0] == "tableend":
            out.append("")
        elif b[0] == "img":
            if b[1] in furniture:
                continue
            ref = f"{assets}/{remap.get(b[1], b[1])}"
            # A space in a markdown image path breaks strict CommonMark.
            # <angle brackets> is the portable form and Typora/Obsidian
            # both accept it, so the NUT1970 "Name.assets" convention
            # survives spaces in the course title.
            out.append(f"![](<{ref}>)" if " " in ref else f"![]({ref})")
            out.append("")
        elif b[0] == "note":
            out.append("> **Notes du présentateur / speaker notes**")
            for l in b[1].splitlines():
                out.append(f"> {l}" if l.strip() else ">")
            out.append("")

    path = os.path.join(outdir, f"{stem}.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out).rstrip() + "\n")
    return path, len(need)


def xlsx_doc(path, max_rows=200, max_cols=40):
    """Spreadsheet -> markdown tables, keeping BOTH the value and the formula.

    Two passes, as the xlsx skill requires: openpyxl with data_only=True gives
    the cached result, data_only=False gives the formula string. For coursework
    this distinction is the whole point -- a GPA calculator's teaching content
    is `=SUM(C2:C9)/COUNT(C2:C9)`, not the 3.4 it happens to display. A static
    table preserves the answer and throws away the method.
    """
    import openpyxl
    doc = Doc(os.path.splitext(os.path.basename(path))[0])
    try:
        wb_v = openpyxl.load_workbook(path, data_only=True, read_only=True)
        wb_f = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception as e:
        doc.stats["open_failed"] = 1
        doc.blocks.append(("p", f"could not open: {e}"))
        return doc

    for name in wb_v.sheetnames:
        sv, sf = wb_v[name], wb_f[name]
        rows_v = list(sv.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True))
        rows_f = list(sf.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True))
        # trim fully-empty trailing rows/cols so a 1048576-row sheet is not a wall
        while rows_v and all(c is None for c in rows_v[-1]):
            rows_v.pop()
            if rows_f: rows_f.pop()
        if not rows_v:
            continue
        width = max((max((i + 1 for i, c in enumerate(r) if c is not None), default=0)
                     for r in rows_v), default=0)
        if width == 0:
            continue
        doc.blocks.append(("h", 2, name))
        doc.stats["sheets"] += 1

        def cell(c):
            return "" if c is None else str(c).replace("|", "\\|").replace("\n", " ")

        def row_cells(ri, r):
            """Prefer the cached value; fall back to the formula when there is
            none. A workbook written by a script and never opened in Excel has
            no cached values at all, and a table of empty cells looks like data
            loss rather than an un-recalculated file."""
            fr = rows_f[ri] if ri < len(rows_f) else ()
            out_ = []
            for ci in range(width):
                v = r[ci] if ci < len(r) else None
                if v is None and ci < len(fr) and isinstance(fr[ci], str) and fr[ci].startswith("="):
                    out_.append("`" + cell(fr[ci]) + "`")
                    doc.stats["uncached_formula_cells"] += 1
                else:
                    out_.append(cell(v))
            return out_

        doc.blocks.append(("tablerow", row_cells(0, rows_v[0])))
        doc.blocks.append(("tablesep", [""] * width))
        for ri, r in enumerate(rows_v[1:], start=1):
            doc.blocks.append(("tablerow", row_cells(ri, r)))
        doc.blocks.append(("tableend", []))
        doc.stats["rows"] += len(rows_v)

        # the formulas, listed once beneath the table rather than inline --
        # inline doubles every column and makes the table unreadable
        formulas = []
        for ri, r in enumerate(rows_f, start=1):
            for ci, c in enumerate(r[:width], start=1):
                if isinstance(c, str) and c.startswith("="):
                    ref = f"{openpyxl.utils.get_column_letter(ci)}{ri}"
                    formulas.append((ref, c))
        if formulas:
            doc.stats["formulas"] += len(formulas)
            doc.blocks.append(("p", "**Formulas in this sheet**"))
            seen = set()
            for ref, fx in formulas:
                # collapse a filled-down column to one representative entry
                shape = re.sub(r"\d+", "#", fx)
                if shape in seen:
                    continue
                seen.add(shape)
                doc.blocks.append(("li", f"`{ref}`  `{fx}`"))
            doc.stats["formula_shapes"] += len(seen)
    wb_v.close(); wb_f.close()
    return doc


def ods_doc(path):
    """OpenDocument spreadsheet via odfpy. Same shape as xlsx, minus formulas:
    odfpy exposes them inconsistently across producers, so rather than emit
    something half-right this reports values only and says so."""
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    doc = Doc(os.path.splitext(os.path.basename(path))[0])
    d = load(path)
    for t in d.spreadsheet.getElementsByType(Table):
        name = t.getAttribute("name") or "Sheet"
        rows = []
        for tr in t.getElementsByType(TableRow):
            cells = []
            for tc in tr.getElementsByType(TableCell):
                rep = int(tc.getAttribute("numbercolumnsrepeated") or 1)
                txt = " ".join(str(p) for p in tc.getElementsByType(P))
                cells.extend([txt] * min(rep, 40))
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(cells)
        if not rows:
            continue
        doc.blocks.append(("h", 2, name))
        w = max(len(r) for r in rows)
        doc.blocks.append(("tablerow", (rows[0] + [""] * w)[:w]))
        doc.blocks.append(("tablesep", [""] * w))
        for r in rows[1:]:
            doc.blocks.append(("tablerow", (r + [""] * w)[:w]))
        doc.stats["sheets"] += 1
        doc.stats["rows"] += len(rows)
    doc.stats["formulas_not_extracted"] = 1
    return doc


def legacy_doc(path):
    """.doc / .xls: convert with LibreOffice first, then read normally.
    Bare `soffice` hangs in some sandboxes; the skill ships a wrapper, so use
    it when present and fall back to soffice otherwise."""
    import tempfile
    out = tempfile.mkdtemp()
    ext = os.path.splitext(path)[1].lower()
    target = "docx" if ext == ".doc" else "xlsx"
    wrapper = "/mnt/skills/public/docx/scripts/office/soffice.py"
    cmd = (["python3", wrapper] if os.path.exists(wrapper) else ["soffice"]) + \
          ["--headless", "--convert-to", target, "--outdir", out, path]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    cand = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + "." + target)
    if not os.path.exists(cand):
        d = Doc(os.path.splitext(os.path.basename(path))[0])
        d.stats["convert_failed"] = 1
        d.blocks.append(("p", f"LibreOffice could not convert: {r.stderr.strip()[:200]}"))
        return d
    d = (docx_doc if target == "docx" else xlsx_doc)(cand)
    d.title = os.path.splitext(os.path.basename(path))[0]
    d.stats["converted_from"] = ext
    shutil.rmtree(out, ignore_errors=True)
    return d


READERS = {".docx": docx_doc, ".pptx": pptx_doc, ".pdf": pdf_doc,
           ".xlsx": xlsx_doc, ".xlsm": xlsx_doc, ".ods": ods_doc,
           ".doc": legacy_doc, ".xls": legacy_doc}


def gather(inputs):
    files = []
    for p in inputs:
        if os.path.isdir(p):
            for dp, dn, fn in os.walk(p):
                dn[:] = [d for d in dn if not d.startswith(".")]
                for f in sorted(fn):
                    if os.path.splitext(f)[1].lower() in READERS and not f.startswith("~$"):
                        files.append(os.path.join(dp, f))
        elif os.path.splitext(p)[1].lower() in READERS:
            files.append(p)
    return files


def main():
    ap = argparse.ArgumentParser(prog="coursekit")
    s = ap.add_subparsers(dest="cmd", required=True)
    c = s.add_parser("convert"); c.add_argument("inputs", nargs="+")
    c.add_argument("-o", default="notes")
    c.add_argument("--drop-repeats", type=int, default=0,
                   metavar="N", help="drop images appearing N+ times (logos, dividers)")
    c.add_argument("--link", action="store_true", help="[[wikilink]] other note titles")
    c.add_argument("--images", choices=["full", "webp", "thumb", "none"], default="full",
                   help="full: as extracted. webp: ~30%% the size, same at reading size. "
                        "thumb: 480px, ~4%%. none: text only -- links stay valid, drop the "
                        "assets folder in later and nothing needs rewriting.")
    p = s.add_parser("probe"); p.add_argument("inputs", nargs="+")
    a = ap.parse_args()

    files = gather(a.inputs)
    if not files:
        sys.exit("no .docx/.pptx/.pdf found")

    docs = []
    for f in files:
        ext = os.path.splitext(f)[1].lower()
        try:
            docs.append((f, READERS[ext](f)))
        except Exception as e:
            print(f"  FAILED {os.path.basename(f)}: {type(e).__name__}: {e}", file=sys.stderr)

    if a.cmd == "probe":
        for f, d in docs:
            hs = sum(1 for b in d.blocks if b[0] == "h")
            ps = sum(1 for b in d.blocks if b[0] == "p")
            ims = sum(1 for b in d.blocks if b[0] == "img")
            print(f"\n{os.path.basename(f)}")
            print(f"   {hs} headings, {ps} paragraphs, {ims} image refs "
                  f"-> {len(d.images)} unique files")
            if d.stats:
                print("   " + "  ".join(f"{k}={v}" for k, v in sorted(d.stats.items())))
        return

    # Two sources can share a stem (Cours3.pptx and Cours3.pdf). Without
    # this the second silently overwrites the first, assets and all.
    stems = Counter(slug(d.title) for _f, d in docs)
    for f, d in docs:
        if stems[slug(d.title)] > 1:
            d.title = f"{d.title} ({os.path.splitext(f)[1].lstrip('.')})"

    titles = [d.title for _f, d in docs]
    total_img = total_uniq = 0
    for f, d in docs:
        path, n = emit(d, a.o, a.drop_repeats, titles if a.link else (), a.images)
        refs = sum(1 for b in d.blocks if b[0] == "img")
        total_img += refs
        total_uniq += n
        extra = f"  [{'  '.join(f'{k}={v}' for k, v in sorted(d.stats.items()))}]" if d.stats else ""
        print(f"{os.path.basename(path):<52} {refs:>4} refs -> {n:>3} files{extra}")
    print(f"\n{len(docs)} notes -> {a.o}/   "
          f"{total_img} image references collapsed to {total_uniq} files")


if __name__ == "__main__":
    main()
